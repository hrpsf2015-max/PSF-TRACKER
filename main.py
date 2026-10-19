import io
import logging
import os
import secrets
from collections import defaultdict
from contextlib import asynccontextmanager
from datetime import date
from decimal import Decimal
from pathlib import Path
from typing import Optional

from fastapi import Depends, FastAPI, HTTPException, Query, Request
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from openpyxl import Workbook
from openpyxl.styles import Font
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.orm import Session, joinedload

from .db import Base, SessionLocal, engine, get_db
from .models import Client, Invoice, Payment, User
from .security import (admin_only, check_lockout, clear_failures, current_user, hash_password,
                       make_token, record_failure, verify_password)

log = logging.getLogger("psf")
PREFIX = os.getenv("INVOICE_PREFIX", "PSF")
STATIC = Path(__file__).parent / "static"
D0 = Decimal("0")
CENT = Decimal("0.01")


@asynccontextmanager
async def lifespan(app: FastAPI):
    Base.metadata.create_all(engine)
    with SessionLocal() as db:
        if not db.scalar(select(User).limit(1)):
            name = os.getenv("ADMIN_USERNAME", "admin")
            pw = os.getenv("ADMIN_PASSWORD")
            if not pw:
                pw = secrets.token_urlsafe(10)
                log.warning("FIRST RUN - admin login created. username=%s password=%s (change it after login)", name, pw)
            db.add(User(username=name, password_hash=hash_password(pw), role="admin"))
            db.commit()
    yield


app = FastAPI(title="PSF Payment Tracker", lifespan=lifespan, docs_url=None, redoc_url=None, openapi_url=None)


@app.middleware("http")
async def security_headers(request: Request, call_next):
    resp = await call_next(request)
    resp.headers["X-Content-Type-Options"] = "nosniff"
    resp.headers["X-Frame-Options"] = "DENY"
    resp.headers["Referrer-Policy"] = "no-referrer"
    resp.headers["Content-Security-Policy"] = (
        "default-src 'self'; script-src 'self' 'unsafe-inline'; style-src 'self' 'unsafe-inline'; img-src 'self' data:"
    )
    if request.url.path.startswith("/api"):
        resp.headers["Cache-Control"] = "no-store"
    return resp


# ---------------------------------------------------------------- helpers
def fy_of(d: date) -> str:
    y = d.year if d.month >= 4 else d.year - 1
    return f"{y}-{str(y + 1)[2:]}"


def net_of(i: Invoice) -> Decimal:
    gst = D0 if i.rcm else (i.gst or D0)
    return (i.gross or D0) + gst - (i.tds or D0)


def paid_of(i: Invoice) -> Decimal:
    return sum((p.amount for p in i.payments), D0)


def status_of(i: Invoice) -> str:
    net, paid = net_of(i), paid_of(i)
    if paid >= net and net >= 0:
        return "Paid"
    if i.due_date and i.due_date < date.today():
        return "Overdue"
    return "Partial" if paid > 0 else "Pending"


def f(x) -> float:
    return float(Decimal(x or 0).quantize(CENT))


def inv_out(i: Invoice, detail: bool = False) -> dict:
    net, paid = net_of(i), paid_of(i)
    out = {
        "id": i.id, "number": i.number, "financial_year": i.financial_year,
        "client_id": i.client_id, "client_name": i.client.name,
        "invoice_date": i.invoice_date.isoformat(),
        "due_date": i.due_date.isoformat() if i.due_date else None,
        "gross": f(i.gross), "gst": f(D0 if i.rcm else i.gst), "rcm": i.rcm, "tds": f(i.tds),
        "net": f(net), "paid": f(paid), "balance": f(net - paid), "status": status_of(i),
        "notes": i.notes,
    }
    if detail:
        out["payments"] = [pay_out(p) for p in i.payments]
    return out


def pay_out(p: Payment) -> dict:
    return {"id": p.id, "invoice_id": p.invoice_id, "payment_date": p.payment_date.isoformat(),
            "amount": f(p.amount), "mode": p.mode, "reference": p.reference, "notes": p.notes}


def load_invoices(db: Session, fy=None, d_from=None, d_to=None, client_id=None):
    q = select(Invoice).options(joinedload(Invoice.client), joinedload(Invoice.payments))
    if fy:
        q = q.where(Invoice.financial_year == fy)
    if d_from:
        q = q.where(Invoice.invoice_date >= d_from)
    if d_to:
        q = q.where(Invoice.invoice_date <= d_to)
    if client_id:
        q = q.where(Invoice.client_id == client_id)
    return list(db.scalars(q.order_by(Invoice.invoice_date.desc(), Invoice.id.desc())).unique())


def check_amounts(gross, gst, tds, rcm):
    gst = D0 if rcm else gst
    if gross + gst - tds < 0:
        raise HTTPException(400, "TDS cannot be more than Gross + GST")


# ---------------------------------------------------------------- schemas
class LoginIn(BaseModel):
    username: str
    password: str


class UserIn(BaseModel):
    username: str = Field(min_length=3, max_length=50, pattern=r"^[A-Za-z0-9_.-]+$")
    password: str = Field(min_length=8, max_length=100)
    role: str = Field(pattern="^(admin|accounts)$")


class UserPatch(BaseModel):
    active: Optional[bool] = None
    role: Optional[str] = Field(default=None, pattern="^(admin|accounts)$")
    password: Optional[str] = Field(default=None, min_length=8, max_length=100)


class PasswordChange(BaseModel):
    old_password: str
    new_password: str = Field(min_length=8, max_length=100)


class ClientIn(BaseModel):
    name: str = Field(min_length=1, max_length=200)
    gstin: Optional[str] = None
    contact_person: Optional[str] = None
    email: Optional[str] = None
    phone: Optional[str] = None
    address: Optional[str] = None
    notes: Optional[str] = None


class InvoiceIn(BaseModel):
    number: str = Field(min_length=1, max_length=60)
    client_id: int
    invoice_date: date
    due_date: Optional[date] = None
    gross: Decimal = Field(ge=0, max_digits=14, decimal_places=2)
    gst: Decimal = Field(default=D0, ge=0, max_digits=14, decimal_places=2)
    rcm: bool = False
    tds: Decimal = Field(default=D0, ge=0, max_digits=14, decimal_places=2)
    notes: Optional[str] = None


class PaymentIn(BaseModel):
    payment_date: date
    amount: Decimal = Field(gt=0, max_digits=14, decimal_places=2)
    mode: Optional[str] = None
    reference: Optional[str] = None
    notes: Optional[str] = None


# ---------------------------------------------------------------- auth & users
@app.post("/api/auth/login")
def login(body: LoginIn, request: Request, db: Session = Depends(get_db)):
    check_lockout(request, body.username)
    user = db.scalar(select(User).where(User.username == body.username.strip()))
    if not user or not user.active or not verify_password(body.password, user.password_hash):
        record_failure(request, body.username)
        raise HTTPException(401, "Wrong username or password")
    clear_failures(request, body.username)
    return {"token": make_token(user), "user": {"username": user.username, "role": user.role}}


@app.get("/api/auth/me")
def me(user: User = Depends(current_user)):
    return {"username": user.username, "role": user.role}


@app.post("/api/auth/change-password")
def change_password(body: PasswordChange, user: User = Depends(current_user), db: Session = Depends(get_db)):
    if not verify_password(body.old_password, user.password_hash):
        raise HTTPException(400, "Current password is wrong")
    user.password_hash = hash_password(body.new_password)
    db.commit()
    return {"ok": True}


@app.get("/api/users")
def list_users(_: User = Depends(admin_only), db: Session = Depends(get_db)):
    return [{"id": u.id, "username": u.username, "role": u.role, "active": u.active}
            for u in db.scalars(select(User).order_by(User.username))]


@app.post("/api/users")
def create_user(body: UserIn, _: User = Depends(admin_only), db: Session = Depends(get_db)):
    if db.scalar(select(User).where(User.username == body.username)):
        raise HTTPException(400, "Username already exists")
    u = User(username=body.username, password_hash=hash_password(body.password), role=body.role)
    db.add(u)
    db.commit()
    return {"id": u.id}


@app.patch("/api/users/{uid}")
def patch_user(uid: int, body: UserPatch, admin: User = Depends(admin_only), db: Session = Depends(get_db)):
    u = db.get(User, uid)
    if not u:
        raise HTTPException(404, "User not found")
    if u.id == admin.id and (body.active is False or (body.role and body.role != "admin")):
        raise HTTPException(400, "You cannot disable or demote your own account")
    if body.active is not None:
        u.active = body.active
    if body.role:
        u.role = body.role
    if body.password:
        u.password_hash = hash_password(body.password)
    db.commit()
    return {"ok": True}


# ---------------------------------------------------------------- clients
def client_out(c: Client, db_rows=None) -> dict:
    out = {k: getattr(c, k) for k in ("id", "name", "gstin", "contact_person", "email", "phone", "address", "notes")}
    invs = c.invoices
    out["invoice_count"] = len(invs)
    out["outstanding"] = f(sum((net_of(i) - paid_of(i) for i in invs), D0))
    return out


@app.get("/api/clients")
def list_clients(_: User = Depends(current_user), db: Session = Depends(get_db)):
    q = select(Client).options(joinedload(Client.invoices).joinedload(Invoice.payments)).order_by(Client.name)
    return [client_out(c) for c in db.scalars(q).unique()]


@app.post("/api/clients")
def create_client(body: ClientIn, _: User = Depends(current_user), db: Session = Depends(get_db)):
    if db.scalar(select(Client).where(Client.name == body.name.strip())):
        raise HTTPException(400, "A client with this name already exists")
    c = Client(**{**body.model_dump(), "name": body.name.strip()})
    db.add(c)
    db.commit()
    return {"id": c.id}


@app.put("/api/clients/{cid}")
def update_client(cid: int, body: ClientIn, _: User = Depends(current_user), db: Session = Depends(get_db)):
    c = db.get(Client, cid)
    if not c:
        raise HTTPException(404, "Client not found")
    clash = db.scalar(select(Client).where(Client.name == body.name.strip(), Client.id != cid))
    if clash:
        raise HTTPException(400, "A client with this name already exists")
    for k, v in body.model_dump().items():
        setattr(c, k, v.strip() if k == "name" else v)
    db.commit()
    return {"ok": True}


@app.delete("/api/clients/{cid}")
def delete_client(cid: int, _: User = Depends(admin_only), db: Session = Depends(get_db)):
    c = db.get(Client, cid)
    if not c:
        raise HTTPException(404, "Client not found")
    if c.invoices:
        raise HTTPException(400, "This client has invoices and cannot be deleted")
    db.delete(c)
    db.commit()
    return {"ok": True}


# ---------------------------------------------------------------- invoices
@app.get("/api/meta")
def meta(_: User = Depends(current_user), db: Session = Depends(get_db)):
    fys = {r for r in db.scalars(select(Invoice.financial_year).distinct())}
    cur = fy_of(date.today())
    fys.add(cur)
    return {"financial_years": sorted(fys, reverse=True), "current_fy": cur}


@app.get("/api/invoices/next-number")
def next_number(on: date, _: User = Depends(current_user), db: Session = Depends(get_db)):
    fy = fy_of(on)
    existing = set(db.scalars(select(Invoice.number).where(Invoice.financial_year == fy)))
    n = len(existing) + 1
    while f"{PREFIX}/{fy}/{n:03d}" in existing:
        n += 1
    return {"number": f"{PREFIX}/{fy}/{n:03d}", "financial_year": fy}


@app.get("/api/invoices")
def list_invoices(
    fy: Optional[str] = None, status: Optional[str] = None, client_id: Optional[int] = None,
    q: Optional[str] = None, _: User = Depends(current_user), db: Session = Depends(get_db),
):
    rows = [inv_out(i) for i in load_invoices(db, fy=fy, client_id=client_id)]
    if status:
        rows = [r for r in rows if r["status"] == status]
    if q:
        ql = q.lower()
        rows = [r for r in rows if ql in r["number"].lower() or ql in r["client_name"].lower()]
    return rows


@app.get("/api/invoices/{iid}")
def get_invoice(iid: int, _: User = Depends(current_user), db: Session = Depends(get_db)):
    i = db.get(Invoice, iid)
    if not i:
        raise HTTPException(404, "Invoice not found")
    return inv_out(i, detail=True)


def _apply_invoice(i: Invoice, body: InvoiceIn, db: Session):
    if not db.get(Client, body.client_id):
        raise HTTPException(400, "Choose a valid client")
    check_amounts(body.gross, body.gst, body.tds, body.rcm)
    fy = fy_of(body.invoice_date)
    number = body.number.strip()
    clash = db.scalar(select(Invoice).where(Invoice.number == number, Invoice.financial_year == fy,
                                            Invoice.id != (i.id or 0)))
    if clash:
        raise HTTPException(400, f"Invoice number {number} already exists in FY {fy}")
    if body.due_date and body.due_date < body.invoice_date:
        raise HTTPException(400, "Due date cannot be before the invoice date")
    i.number, i.financial_year, i.client_id = number, fy, body.client_id
    i.invoice_date, i.due_date = body.invoice_date, body.due_date
    i.gross, i.rcm, i.tds, i.notes = body.gross, body.rcm, body.tds, body.notes
    i.gst = D0 if body.rcm else body.gst


@app.post("/api/invoices")
def create_invoice(body: InvoiceIn, user: User = Depends(current_user), db: Session = Depends(get_db)):
    i = Invoice(created_by=user.username)
    _apply_invoice(i, body, db)
    db.add(i)
    db.commit()
    return {"id": i.id}


@app.put("/api/invoices/{iid}")
def update_invoice(iid: int, body: InvoiceIn, _: User = Depends(current_user), db: Session = Depends(get_db)):
    i = db.get(Invoice, iid)
    if not i:
        raise HTTPException(404, "Invoice not found")
    _apply_invoice(i, body, db)
    if net_of(i) < paid_of(i):
        db.rollback()
        raise HTTPException(400, "Net receivable cannot be less than payments already received")
    db.commit()
    return {"ok": True}


@app.delete("/api/invoices/{iid}")
def delete_invoice(iid: int, _: User = Depends(admin_only), db: Session = Depends(get_db)):
    i = db.get(Invoice, iid)
    if not i:
        raise HTTPException(404, "Invoice not found")
    db.delete(i)
    db.commit()
    return {"ok": True}


@app.post("/api/invoices/{iid}/payments")
def add_payment(iid: int, body: PaymentIn, user: User = Depends(current_user), db: Session = Depends(get_db)):
    i = db.get(Invoice, iid)
    if not i:
        raise HTTPException(404, "Invoice not found")
    balance = net_of(i) - paid_of(i)
    if body.amount > balance:
        raise HTTPException(400, f"Amount is more than the balance of {balance}")
    db.add(Payment(invoice_id=i.id, created_by=user.username, **body.model_dump()))
    db.commit()
    return {"ok": True}


@app.delete("/api/payments/{pid}")
def delete_payment(pid: int, _: User = Depends(admin_only), db: Session = Depends(get_db)):
    p = db.get(Payment, pid)
    if not p:
        raise HTTPException(404, "Payment not found")
    db.delete(p)
    db.commit()
    return {"ok": True}


# ---------------------------------------------------------------- dashboard
@app.get("/api/dashboard")
def dashboard(fy: Optional[str] = None, _: User = Depends(current_user), db: Session = Depends(get_db)):
    invs = load_invoices(db, fy=fy)
    rows = [inv_out(i) for i in invs]
    tot = lambda k: round(sum(r[k] for r in rows), 2)
    by_status = defaultdict(lambda: {"count": 0, "balance": 0.0})
    by_client = defaultdict(float)
    for r in rows:
        s = by_status[r["status"]]
        s["count"] += 1
        s["balance"] = round(s["balance"] + r["balance"], 2)
        if r["balance"] > 0:
            by_client[r["client_name"]] += r["balance"]
    top = sorted(by_client.items(), key=lambda x: -x[1])[:5]
    return {
        "invoice_count": len(rows), "gross": tot("gross"), "gst": tot("gst"), "tds": tot("tds"),
        "net": tot("net"), "received": tot("paid"), "outstanding": tot("balance"),
        "by_status": {k: by_status[k] for k in ("Pending", "Partial", "Overdue", "Paid") if k in by_status},
        "top_outstanding": [{"client": c, "balance": round(b, 2)} for c, b in top],
    }


# ---------------------------------------------------------------- reports & export
REPORTS = {"monthly": "Monthly", "daily": "Daily collections", "client": "Client-wise", "tds": "TDS",
           "gst": "GST / RCM", "overdue": "Overdue"}


def build_report(db: Session, kind: str, fy, d_from, d_to):
    if kind not in REPORTS:
        raise HTTPException(404, "Unknown report")
    if kind == "daily":
        q = select(Payment).options(joinedload(Payment.invoice).joinedload(Invoice.client))
        if fy:
            q = q.join(Invoice).where(Invoice.financial_year == fy)
        if d_from:
            q = q.where(Payment.payment_date >= d_from)
        if d_to:
            q = q.where(Payment.payment_date <= d_to)
        pays = db.scalars(q.order_by(Payment.payment_date.desc())).unique()
        cols = ["Date", "Invoice", "Client", "Amount", "Mode", "Reference"]
        rows = [[p.payment_date.isoformat(), p.invoice.number, p.invoice.client.name, f(p.amount),
                 p.mode or "", p.reference or ""] for p in pays]
        return cols, rows
    invs = [inv_out(i) for i in load_invoices(db, fy=fy, d_from=d_from, d_to=d_to)]
    if kind in ("monthly", "client"):
        groups = defaultdict(lambda: [0, 0.0, 0.0, 0.0, 0.0, 0.0, 0.0])
        for r in invs:
            key = r["invoice_date"][:7] if kind == "monthly" else r["client_name"]
            g = groups[key]
            g[0] += 1
            for n, k in enumerate(("gross", "gst", "tds", "net", "paid", "balance"), 1):
                g[n] += r[k]
        cols = ["Month" if kind == "monthly" else "Client", "Invoices", "Gross", "GST", "TDS",
                "Net receivable", "Received", "Outstanding"]
        rows = [[k] + [g[0]] + [round(x, 2) for x in g[1:]] for k, g in sorted(groups.items())]
        return cols, rows
    if kind == "tds":
        cols = ["Invoice", "Client", "Date", "FY", "Gross", "TDS"]
        return cols, [[r["number"], r["client_name"], r["invoice_date"], r["financial_year"], r["gross"], r["tds"]]
                      for r in invs if r["tds"] > 0]
    if kind == "gst":
        cols = ["Invoice", "Client", "Date", "FY", "Gross", "GST billed", "Reverse charge"]
        return cols, [[r["number"], r["client_name"], r["invoice_date"], r["financial_year"], r["gross"], r["gst"],
                       "Yes" if r["rcm"] else "No"] for r in invs if r["gst"] > 0 or r["rcm"]]
    cols = ["Invoice", "Client", "Due date", "Days overdue", "Balance"]
    today = date.today()
    return cols, [[r["number"], r["client_name"], r["due_date"],
                   (today - date.fromisoformat(r["due_date"])).days, r["balance"]]
                  for r in invs if r["status"] == "Overdue"]


def _sheet(wb: Workbook, title: str, cols: list, rows: list, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title[:31]
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    for n, col in enumerate(cols, 1):
        width = max([len(str(col))] + [len(str(r[n - 1])) for r in rows[:200]]) + 2
        ws.column_dimensions[ws.cell(1, n).column_letter].width = min(width, 45)
    ws.freeze_panes = "A2"


def _xlsx(wb: Workbook, name: str):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{name}"'})


@app.get("/api/reports/{kind}")
def report(kind: str, fy: Optional[str] = None, date_from: Optional[date] = None,
           date_to: Optional[date] = None, _: User = Depends(current_user), db: Session = Depends(get_db)):
    cols, rows = build_report(db, kind, fy, date_from, date_to)
    return {"title": REPORTS[kind], "columns": cols, "rows": rows}


@app.get("/api/reports/{kind}/excel")
def report_excel(kind: str, fy: Optional[str] = None, date_from: Optional[date] = None,
                 date_to: Optional[date] = None, _: User = Depends(current_user), db: Session = Depends(get_db)):
    cols, rows = build_report(db, kind, fy, date_from, date_to)
    wb = Workbook()
    _sheet(wb, REPORTS[kind], cols, rows, first=True)
    return _xlsx(wb, f"report-{kind}-{date.today()}.xlsx")


@app.get("/api/export/all")
def export_all(_: User = Depends(current_user), db: Session = Depends(get_db)):
    wb = Workbook()
    invs = load_invoices(db)
    cols = ["Invoice no.", "FY", "Client", "Invoice date", "Due date", "Gross", "GST", "RCM", "TDS",
            "Net receivable", "Received", "Balance", "Status", "Notes"]
    rows = []
    for i in invs:
        r = inv_out(i)
        rows.append([r["number"], r["financial_year"], r["client_name"], r["invoice_date"], r["due_date"] or "",
                     r["gross"], r["gst"], "Yes" if r["rcm"] else "No", r["tds"], r["net"], r["paid"],
                     r["balance"], r["status"], r["notes"] or ""])
    _sheet(wb, "Invoices", cols, rows, first=True)
    pcols = ["Invoice no.", "Client", "Payment date", "Amount", "Mode", "Reference", "Notes"]
    prow = [[i.number, i.client.name, p.payment_date.isoformat(), f(p.amount), p.mode or "", p.reference or "",
             p.notes or ""] for i in invs for p in i.payments]
    _sheet(wb, "Payments", pcols, prow)
    clients = db.scalars(select(Client).order_by(Client.name)).all()
    _sheet(wb, "Clients", ["Name", "GSTIN", "Contact", "Email", "Phone", "Address", "Notes"],
           [[c.name, c.gstin or "", c.contact_person or "", c.email or "", c.phone or "", c.address or "",
             c.notes or ""] for c in clients])
    return _xlsx(wb, f"psf-all-data-{date.today()}.xlsx")


# ---------------------------------------------------------------- front end
@app.get("/", include_in_schema=False)
def index():
    return FileResponse(STATIC / "index.html")


@app.get("/healthz", include_in_schema=False)
def health():
    return {"ok": True}
